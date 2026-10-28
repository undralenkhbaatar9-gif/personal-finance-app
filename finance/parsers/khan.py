
import csv
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from finance.parsers.base import (
    BankParser,
    FilePath,
    StatementInfo,
    StdTxn,
    to_decimal,
)

IBAN_PATTERN = re.compile(r"MN\d{18}", re.IGNORECASE)
DATE_PATTERN = re.compile(r"\d{4}/\d{2}/\d{2}")


def _cell_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    match = DATE_PATTERN.search(_cell_text(value))

    if match is None:
        return None

    return datetime.strptime(match.group(), "%Y/%m/%d").date()


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    value_text = _cell_text(value)

    for date_format in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(value_text, date_format)
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(value_text)
    except ValueError:
        return None


def _optional_decimal(value: Any) -> Decimal | None:
    if value is None or _cell_text(value) == "":
        return None

    return to_decimal(value)


class KhanParser(BankParser):
    code = "khan"
    name = "Khan Bank"

    def detect(self, path: FilePath) -> bool:
        file_path = Path(path)

        if file_path.suffix.lower() == ".csv":
            try:
                with file_path.open(
                    encoding="utf-8-sig",
                    newline="",
                ) as statement_file:
                    row = next(csv.reader(statement_file), [])

                text = " ".join(
                    _cell_text(value) for value in row
                ).upper()

                match = IBAN_PATTERN.search(text)

                if "IBAN" not in text or match is None:
                    return False

                digits = match.group()[2:]
                return digits[2:6] == "0005"

            except (OSError, UnicodeError, csv.Error):
                return False

        try:
            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )

            try:
                worksheet = workbook.worksheets[0]

                row = next(
                    worksheet.iter_rows(
                        min_row=1,
                        max_row=1,
                        values_only=True,
                    ),
                    (),
                )

                text = " ".join(
                    _cell_text(value) for value in row
                ).upper()

                match = IBAN_PATTERN.search(text)

                if "IBAN" not in text or match is None:
                    return False

                digits = match.group()[2:]
                return digits[2:6] == "0005"

            finally:
                workbook.close()

        except (OSError, ValueError, IndexError, StopIteration):
            return False

    def parse(
        self,
        path: FilePath,
    ) -> tuple[StatementInfo, list[StdTxn]]:
        file_path = Path(path)

        if file_path.suffix.lower() == ".csv":
            try:
                with file_path.open(
                    encoding="utf-8-sig",
                    newline="",
                ) as statement_file:
                    rows = list(csv.reader(statement_file))
            except (OSError, UnicodeError, csv.Error) as error:
                raise ValueError(
                    f"Could not read Khan CSV statement: {error}"
                ) from error

            if not rows:
                raise ValueError("Khan CSV statement is empty")

            return self._parse_rows(
                rows[0],
                iter(rows[2:]),
            )

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.worksheets[0]

            info_row = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )

            data_rows = worksheet.iter_rows(
                min_row=2,
                values_only=True,
            )

            return self._parse_rows(info_row, data_rows)

        finally:
            workbook.close()

    def _parse_rows(
        self,
        info_row: Any,
        data_rows: Any,
    ) -> tuple[StatementInfo, list[StdTxn]]:
        info_text = " ".join(
            _cell_text(value) for value in info_row
        )

        iban_match = IBAN_PATTERN.search(info_text)

        if iban_match is None:
            raise ValueError(
                "Khan statement row 1 does not contain a valid MN IBAN"
            )

        iban = iban_match.group().upper()
        account_no = iban[-10:]

        dates = [
            parsed
            for value in info_row
            if (parsed := _parse_date(value)) is not None
        ]

        currency = next(
            (
                _cell_text(value).upper()
                for value in info_row
                if _cell_text(value).upper() == "MNT"
            ),
            "MNT",
        )

        transactions: list[StdTxn] = []
        total_credit: Decimal | None = None
        total_debit: Decimal | None = None

        for row in data_rows:
            values = list(row)
            values += [None] * max(0, 8 - len(values))

            raw_datetime = values[0]
            row_text = " ".join(
                _cell_text(value) for value in values
            ).casefold()

            if "нийт дүн" in row_text or row_text.strip().startswith("нийт"):
                if values[3] not in (None, ""):
                    total_credit = abs(to_decimal(values[3]))

                if values[4] not in (None, ""):
                    total_debit = abs(to_decimal(values[4]))

                continue

            if raw_datetime is None or _cell_text(raw_datetime) == "":
                continue

            txn_datetime = _parse_datetime(raw_datetime)

            if txn_datetime is None:
                continue

            credit_value = _optional_decimal(values[3])
            debit_value = _optional_decimal(values[4])

            credit = abs(credit_value or Decimal("0.00"))
            debit = abs(debit_value or Decimal("0.00"))

            if credit != 0 and debit != 0:
                raise ValueError(
                    "A transaction row contains both credit and debit "
                    f"amounts: {txn_datetime}"
                )

            amount = credit - debit

            balance_before = _optional_decimal(values[2])
            balance_after = _optional_decimal(values[5])

            transactions.append(
                StdTxn(
                    txn_datetime=txn_datetime,
                    amount=amount,
                    balance_before=balance_before,
                    balance_after=balance_after,
                    description=_cell_text(values[6]),
                    counterparty_account=(
                        _cell_text(values[7]) or None
                    ),
                    branch_code=_cell_text(values[1]) or None,
                )
            )

        if not transactions:
            raise ValueError(
                "Khan statement contains no transaction rows"
            )

        info = StatementInfo(
            account_no=account_no,
            iban=iban,
            currency=currency,
            period_from=dates[0] if dates else None,
            period_to=dates[1] if len(dates) > 1 else None,
            total_credit=total_credit,
            total_debit=total_debit,
        )

        return info, transactions
